import io
import base64
from odoo.fields import Binary
from odoo.api import Environment
from odoo.exceptions import UserError
from openpyxl import Workbook, load_workbook
from .models import Template


class TemplateProjectService:

    def __init__(self, env: Environment):
        self.env = env
        self.model = env["pkf.timesheet.project"]

    def _load_template(self, file_data: Binary) -> Workbook:
        file_data = base64.b64decode(file_data)
        file_stream = io.BytesIO(file_data)
        try:
            return load_workbook(file_stream)
        except Exception as e:
            raise UserError(f"El archivo no es un archivo de Excel válido: {e}")

    def _build_template_records(self, wb: Workbook) -> list[Template]:
        sheet = wb.active
        records: list[Template] = []
        row = 2
        while True:
            record = Template.from_wb_sheet(sheet, row)

            if record is None:
                # Si la celda nombre está vacía → fin de datos
                if not sheet.cell(row, 1).value:
                    break
                # Si es inválida → saltar fila
                row += 1
                continue

            records.append(record)
            row += 1

        return records

    def _set_partner(self, values: dict, rec: Template):
        partner_model = self.env["res.partner"]
        partner = partner_model.browse(rec.client_id)

        values["partner_id"] = partner.id if partner else False

        if not partner:
            values["partner_group_id"] = False
        else:
            values["partner_group_id"] = partner.pkf_partner_group.id

    def _set_user_info(self, values: dict, rec: Template):
        employee_model = self.env["hr.employee"]
        employee = employee_model.search(
            [("work_email", "=", rec.owner_email)], limit=1
        )

        values["manager_id"] = employee.id if employee else False

    def _set_phase_group(self, values: dict):
        group_id = self.env["pkf.timesheet.phase.group"].search(
            [("is_all", "=", True)], limit=1
        )
        values["phase_group_id"] = group_id.id if group_id else False

    def _create_projects(self, template_records: list[Template]):
        for rec in template_records:
            values = {"name": rec.name}
            self._set_partner(values, rec)
            self._set_user_info(values, rec)
            self._set_phase_group(values)

            values["total_budget_amount"] = rec.budget
            values["cancellation_porcentage"] = rec.cancelation
            values["state"] = "in_progress"
            values["period_type"] = "biweekly"
            values["note"] = rec.note

            project = self.model.create(values)

            if project:
                project._onchange_period_type()
                project._onchange_manager_id()
                project._onchange_group_phase_id()

    def import_template(self, file_data: Binary) -> bool:
        wb = self._load_template(file_data)
        records = self._build_template_records(wb)
        if not records:
            return False

        self._create_projects(records)
        self.model.refresh_codes()
        return True
